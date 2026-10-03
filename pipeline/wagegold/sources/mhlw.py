"""Ministry of Health, Labour and Welfare of Japan (厚生労働省): its own release of the
series ILOSTAT republishes for Japan as DA:260 - the Basic Survey on Wage Structure
(賃金構造基本統計調査, June each year): scheduled cash earnings (所定内給与額) of
regular employees (一般労働者), both sexes, all industries, enterprises with 10 or
more employees - and, from the same table, their scheduled hours actually worked
(所定内実労働時間数, hours in June).

ILOSTAT republishes it with a delay; MHLW's yearly tables (released each March on
e-Stat) continue it (build.UnitGraph._extend, which first checks that the two agree on
the years both have).  The tables are discovered on e-Stat's file listing of the
survey: each survey year's folder "■令和N年賃金構造基本統計調査" → "一般労働者" →
"産業大分類" holds table 1 (学歴、年齢階級別きまって支給する現金給与額、所定内給与額
及び年間賞与その他特別給与額, 産業計・産業別) as an Excel file.

MHLW changed the survey's estimation method with the 2020 survey and states that its
figures from 2020 on are not comparable with those before, so the series read here
starts with the 2020 survey (FIRST_YEAR).
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from urllib.parse import urljoin

from ..fetch import Fetcher, Snapshot
from ..model import Obs
from .common import check_prefix

PUBLISHER = "mhlw"  # catalog keys src.mhlw.*
SERIES = "ext_ilo_monthly_mean@DA:260"  # continues ILOSTAT's ilo_monthly_mean@DA:260 for JPN
HOURS = "ext_ilo_weekly_hours@DA:260"  # the same survey's hours (ILOSTAT has none for DA:260)
BASE = "https://www.e-stat.go.jp"
LIST_URL = BASE + "/stat-search/files?page=1&toukei=00450091&tstat=000001011429"
FIRST_YEAR = 2020  # the survey's current estimation method (see above)
WEEKS_PER_MONTH = 52 / 12
YEAR_TITLE = re.compile(r"^■令和(\d+|元)年賃金構造基本統計調査$")
# A node of the listing's folder tree: its attributes (data-matter = depth,
# data-valueN = the tclassN ids of its path) and the title that follows.
NODE_RE = re.compile(r"<(?:div|a)\b([^>]*\bdata-matter=\"(\d)\"[^>]*)>")
NODE_TITLE_RE = re.compile(r"class=\"stat-title[^\"]*\"[^>]*>\s*<span[^>]*>\s*([^<]+?)\s*<span class=\"stat-pc\">")
ROW_SPLIT = '<article class="stat-dataset_list-item">'
TABLE_NO_RE = re.compile(r"表番号&nbsp;</span>\s*<span>\s*(\d+)\s*</span>")
SURVEY_YEAR_RE = re.compile(r"調査年月&nbsp;&nbsp;</span>\s*(\d{4})年")
EXCEL_RE = re.compile(r"href=\"(/stat-search/file-download\?statInfId=(\d+)&(?:amp;)?fileKind=\d+)\"[^>]*data-file_type=\"EXCEL")


@dataclass
class Table:
    year: int
    url: str
    key: str


def _text(s: str) -> str:
    return unicodedata.normalize("NFKC", re.sub(r"\s+", " ", s)).strip()


def survey_year(folder_title: str) -> int | None:
    m = YEAR_TITLE.match(_text(folder_title))
    if not m:
        return None
    return 2018 + (1 if m.group(1) == "元" else int(m.group(1)))


def folders(html: str) -> list[tuple[list[str], str, str | None]]:
    """The listing's folder tree: (tclass ids of the path, title, href or None)."""
    out = []
    for m in NODE_RE.finditer(html):
        attrs, depth = m.group(1), int(m.group(2))
        ids = [re.search(rf"data-value{n}=\"(\d+)\"", attrs) for n in range(1, depth + 1)]
        title = NODE_TITLE_RE.search(html, m.end())
        if not all(ids) or not title:
            continue
        href = re.search(r"\bhref=\"([^\"]+)\"", attrs)
        out.append(([i.group(1) for i in ids], _text(title.group(1)), href.group(1).replace("&amp;", "&") if href else None))
    return out


def industry_folders(html: str) -> dict[int, str]:
    """Survey year -> URL of its 一般労働者 / 産業大分類 folder."""
    tree = folders(html)
    title = {tuple(path): t for path, t, _h in tree}
    out: dict[int, str] = {}
    for path, t, href in tree:
        if len(path) != 3 or t != "産業大分類" or not href:
            continue
        year = survey_year(title.get(tuple(path[:1]), ""))
        if year is not None and year >= FIRST_YEAR and title.get(tuple(path[:2])) == "一般労働者":
            out[year] = urljoin(BASE, href)
    return out


def table1(html: str, year: int) -> str:
    """URL of table 1's Excel file on a 産業大分類 folder page."""
    for row in html.split(ROW_SPLIT)[1:]:
        no, yr, xl = TABLE_NO_RE.search(row), SURVEY_YEAR_RE.search(row), EXCEL_RE.search(row)
        if no and no.group(1) == "1" and xl:
            if not yr or int(yr.group(1)) != year:
                raise ValueError(f"table 1 of the {year} folder is for {yr.group(1) if yr else 'no'} survey year")
            return urljoin(BASE, xl.group(1))
    raise ValueError(f"no table 1 Excel file in the {year} folder")


def discover(f: Fetcher) -> list[Table]:
    out = []
    for year, folder in sorted(industry_folders(f.get_transient(LIST_URL).decode("utf-8", "replace")).items()):
        url = table1(f.get_transient(folder).decode("utf-8", "replace"), year)
        out.append(Table(year, url, f"mhlw/bsws/{year}_ippan_sangyo_t1"))
    if not out:
        raise ValueError("no survey year folder found on the e-Stat listing")
    return out


def collect(f: Fetcher) -> list[Obs]:
    snaps: list[tuple[Snapshot, int]] = []
    if not f.offline:
        for t in discover(f):
            snaps.append((f.get(t.key, t.url, ext="xlsx", check=check_prefix(b"PK", "an xlsx workbook")), t.year))
    seen = {s.key for s, _y in snaps}
    snaps += [(s, int(s.key.rsplit("/", 1)[1][:4])) for s in f.committed("mhlw/bsws/") if s.key not in seen]
    out = []
    for snap, year in snaps:
        earnings, hours = parse(snap.read(), snap.key)
        out.append(Obs(SERIES, "JPN", str(year), earnings, snap.key, PUBLISHER))
        out.append(Obs(HOURS, "JPN", str(year), hours / WEEKS_PER_MONTH, snap.key, PUBLISHER))
    return out


def parse(xlsx: bytes, snapshot: str) -> tuple[float, float]:
    """(所定内給与額 in yen, 所定内実労働時間数 in hours a month) of 男女計, 学歴計 (all
    ages), 企業規模計(10人以上), 産業計."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets[:3]:
        lines.append(f"sheet {ws.title!r}")
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i >= 25:
                break
            lines.append(" | ".join("" if c is None else str(c) for c in r)[:300])
    raise ValueError(f"{snapshot}: layout not yet known; sheets {wb.sheetnames[:10]}\n" + "\n".join(lines))
